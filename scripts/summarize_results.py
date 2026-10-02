# The experiment workbook, rebuilt from every run folder in result/.
#
#   python scripts/summarize_results.py              reads result/*, writes into result/
#   python scripts/summarize_results.py --result-dir path/to/result
#
# scripts/run_all.sh calls it after every run, so the workbook always matches
# the folders on disk; the folders (reports, probabilities, run_info.txt) are
# the record, the workbook is a view of them. It is written to a temporary file
# and then renamed, so two runs finishing together cannot corrupt it.
#
# Every metric is recomputed from the saved probabilities (*_validation_probs.npz,
# *_test_probs.npz) at the threshold t* stored with them, so old runs get the
# metrics added later (accuracy, specificity) and all rows are computed the same
# way. A run without probability files falls back to the numbers in its reports.
#
# Output:
#   so_thi_nghiem.xlsx      sheets
#     Bang_chinh            one row per scenario of scripts/scenarios.py (the run tag
#                           is the scenario code): test metrics mean ± std over the
#                           latest run of each seed, Δ AUC and DeLong against M0 and M-any
#     Kich_ban              the scenario catalog: question, script, arguments, bundle
#     Lan_chay              one row per run folder: date, git, host, GPU, every
#                           hyperparameter, val and test metrics mean ± std
#     Theo_seed             one row per seed
#     DeLong                paired DeLong test per seed, every scenario vs M0 and vs M-any
#     Theo_cau_hinh         one row per configuration (run name without the seed),
#                           also for runs made before the scenario codes existed
#     Giai_thich            what the columns mean
#   tong_hop_lan_chay.csv, tong_hop_cau_hinh.csv, tong_hop_theo_seed.csv, tong_hop_bang_chinh.csv
#   tong_hop_ket_qua.md     short Markdown table of Bang_chinh

import argparse
import csv
import json
import os
import re
import statistics
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
from sklearn.metrics import average_precision_score, precision_recall_fscore_support, roc_auc_score

sys.path.insert(0, str(Path(__file__).resolve().parent))
from delong import delong  # noqa: E402
from scenarios import SCENARIOS, hypergraph_file  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = "so_thi_nghiem.xlsx"
# Every scenario is compared with these two (Bang_chinh, DeLong).
REFERENCES = ("M0", "M-any")

# (key, header): metrics computed for validation and test.
METRICS = (
    ("auroc", "AUROC (AUC)"),
    ("auprc", "AUPRC"),
    ("accuracy", "ACC"),
    ("precision", "Precision"),
    ("recall", "Recall"),
    ("f1", "F1"),
    ("specificity", "Specificity"),
    ("macro_f1", "Macro-F1"),
    ("f1_negative", "F1 lớp không bỏ"),
    ("auprc_negative", "AUPRC lớp không bỏ"),
)
# (key in the train settings, header)
HYPERPARAMETERS = (
    ("families", "Hyperedge"),
    ("hypergraph", "File graph"),
    ("causal", "Causal"),
    ("shuffle_graph", "Xáo graph (seed)"),
    ("skip_connection", "Skip"),
    ("family_weights", "W family"),
    ("edge_weights", "W từng hyperedge"),
    ("edge_sampling", "Me"),
    ("node_sampling", "Mv"),
    ("add_per_edge", "ΔH/cạnh"),
    ("lambda_cl", "λ contrastive"),
    ("temperature", "τ contrastive"),
    ("hidden_dim", "Chiều ẩn"),
    ("dropout", "Dropout"),
    ("learning_rate", "LR"),
    ("weight_decay", "Weight decay"),
    ("family_weight_lr", "LR W family"),
    ("edge_weight_lr", "LR W hyperedge"),
    ("epochs", "Epoch tối đa"),
    ("patience", "Patience"),
    ("eval_every", "Validate mỗi"),
    ("validation_limit", "Tập con val"),
    ("select_metric", "Chọn checkpoint"),
    ("pos_weight", "pos_weight"),
    ("lr_schedule", "LR schedule"),
    ("feature_set", "Feature"),
)
# Learned structure at the best epoch (HSL keep rates, W).
STRUCTURE = ("kept_course", "kept_object", "kept_behavioral", "kept_user",
             "w_course", "w_object", "w_behavioral", "w_user", "w_self_loop",
             "alpha_course", "alpha_object", "alpha_behavioral", "alpha_user")


def metrics_at(labels, probabilities, threshold):
    predicted = probabilities >= threshold
    precision, recall, f1, _ = precision_recall_fscore_support(
        labels, predicted, labels=[0, 1], zero_division=0
    )
    return {
        "auroc": roc_auc_score(labels, probabilities),
        "auprc": average_precision_score(labels, probabilities),
        "accuracy": float(np.mean(predicted == labels)),
        "precision": precision[1],
        "recall": recall[1],
        "f1": f1[1],
        "specificity": recall[0],
        "macro_f1": f1.mean(),
        "f1_negative": f1[0],
        "auprc_negative": average_precision_score(1 - labels, 1 - probabilities),
    }


# Metrics of one split from the probability file, else from the report.
def split_metrics(reports, run_name, split, fallback):
    path = reports / f"{run_name}_{split}_probs.npz"
    if path.exists():
        with np.load(path) as saved:
            labels = saved["labels"].astype(int)
            return metrics_at(labels, saved["probabilities"], float(saved["threshold"]))
    fallback = fallback or {}
    return {key: fallback.get("auc" if key == "auroc" else key) for key, _ in METRICS}


def read_run_info(run_dir):
    info = {}
    path = run_dir / "run_info.txt"
    if path.exists():
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            if ":" in line:
                key, value = line.split(":", 1)
                info[key.strip()] = value.strip()
    return info


def read_durations(run_dir):
    path = run_dir / "manifest.tsv"
    if not path.exists():
        return {}
    with open(path, encoding="utf-8") as manifest:
        return {int(row["seed"]): int(row["duration_sec"]) for row in csv.DictReader(manifest, delimiter="\t")}


def folder_time(run_dir):
    try:
        return datetime.strptime(run_dir.name[:16], "%d-%m-%Y_%H-%M")
    except ValueError:
        return datetime.fromtimestamp(run_dir.stat().st_mtime)


# settings["model"] of src/9_baselines.py and of the published baselines (files 11-15).
BASELINE_NAMES = {"gbdt": "GBDT", "logreg": "LR", "hypergcn": "HyperGCN", "signet": "SIG-Net",
                  "mstgcn": "MST-GCN", "catfhn": "CA-TFHN", "cfin": "CFIN"}


def model_name(settings):
    if settings.get("model") in BASELINE_NAMES:
        return BASELINE_NAMES[settings["model"]]
    families = [f for f in settings.get("families", ["course", "object", "behavioral"]) if f != "self_loop"]
    if not families and not settings.get("hsl", True):
        return "MLP"
    return "HSL" if settings.get("hsl", True) else "HGNN"


def seed_rows(run_dir):
    reports = run_dir / "reports"
    info = read_run_info(run_dir)
    durations = read_durations(run_dir)
    rows = []
    for train_path in sorted(reports.glob("*_train.json")):
        run_name = train_path.name.removesuffix("_train.json")
        train = json.loads(train_path.read_text(encoding="utf-8"))
        test_path = reports / f"{run_name}_test.json"
        test = json.loads(test_path.read_text(encoding="utf-8")) if test_path.exists() else {}
        settings = train.get("settings", {})
        best = next((r for r in train.get("history", []) if r.get("epoch") == train.get("best_epoch")), {})
        row = {
            "folder": run_dir.name,
            "started": folder_time(run_dir),
            "note": info.get("note", ""),
            "git": info.get("git", ""),
            "host": info.get("host", ""),
            "gpu": info.get("gpu", ""),
            "python": info.get("python", ""),
            "torch": info.get("torch", ""),
            "scenario": re.sub(r"_seed_\d+$", "", run_name),
            "tag": settings.get("tag", ""),
            # Read from the bundle by 8_train.py; runs before that left it out.
            "user_rule": train.get("user_rule"),
            "test_probs": reports / f"{run_name}_test_probs.npz",
            "model": model_name(settings),
            "seed": train.get("seed"),
            "best_epoch": train.get("best_epoch"),
            "epochs_run": len(train.get("history", [])) or None,
            "threshold": train.get("threshold"),
            "duration_min": round(durations[train["seed"]] / 60, 1) if train.get("seed") in durations else None,
            "settings": settings,
            "structure": {key: best.get(key) for key in STRUCTURE},
        }
        validation = split_metrics(reports, run_name, "validation", train.get("best_validation"))
        tested = split_metrics(reports, run_name, "test", test.get("test"))
        row.update({f"val_{key}": value for key, value in validation.items()})
        row.update({f"test_{key}": value for key, value in tested.items()})
        rows.append(row)
    return rows


def mean_std(values):
    values = [float(v) for v in values if v is not None]
    if not values:
        return None, None
    return statistics.mean(values), (statistics.stdev(values) if len(values) > 1 else 0.0)


def text(value):
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value if item != "self_loop")
    return "" if value is None else value


# One summary row over a group of seed rows (one folder, or one configuration).
def summarize(group, first_columns):
    last = group[-1]
    row = dict(first_columns)
    row["Mô hình"] = last["model"]
    row["Kịch bản (tag)"] = last["tag"]
    row["User rule"] = text(last["user_rule"])
    row["Seeds"] = " ".join(str(r["seed"]) for r in sorted(group, key=lambda r: r["seed"]))
    row["Số seed"] = len(group)
    for key, header in HYPERPARAMETERS:
        row[header] = text(last["settings"].get(key))
    for split, split_name in (("test", "test"), ("val", "val")):
        for key, header in METRICS:
            mean, std = mean_std(r[f"{split}_{key}"] for r in group)
            row[f"{header} {split_name}"] = "" if mean is None else f"{mean:.4f} ± {std:.4f}"
    for column, key in (("Ngưỡng t*", "threshold"), ("Best epoch", "best_epoch"),
                        ("Epoch đã chạy", "epochs_run"), ("Phút/seed", "duration_min")):
        mean, std = mean_std(r[key] for r in group)
        row[column] = "" if mean is None else (f"{mean:.4f}" if key == "threshold" else f"{mean:.0f}")
    for key in STRUCTURE:
        mean, _ = mean_std(r["structure"][key] for r in group)
        row[key] = "" if mean is None else round(mean, 4)
    row["_sort"] = mean_std(r["test_auroc"] for r in group)[0] or 0.0
    return row


def seed_table(rows):
    table = []
    for r in rows:
        row = {"Lần chạy": r["folder"], "Kịch bản (tag)": r["tag"], "Cấu hình": r["scenario"],
               "Mô hình": r["model"], "User rule": text(r["user_rule"]), "Seed": r["seed"],
               "Best epoch": r["best_epoch"],
               "Epoch đã chạy": r["epochs_run"], "Ngưỡng t*": r["threshold"], "Phút": r["duration_min"]}
        for split in ("test", "val"):
            for key, header in METRICS:
                value = r[f"{split}_{key}"]
                row[f"{header} {split}"] = None if value is None else round(float(value), 4)
        row.update({key: r["structure"][key] for key in STRUCTURE})
        table.append(row)
    return table


# Columns in order of first appearance over all rows, so a row that lacks a
# column (a scenario not run yet) does not drop it for the others.
def table_columns(table):
    return list(dict.fromkeys(c for row in table for c in row if not c.startswith("_")))


def write_csv(path, table):
    with open(path, "w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(output, fieldnames=table_columns(table), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(table)


# Latest run of every (value of key, seed): folders are in time order, so a
# later run of a seed replaces an earlier one, and seeds run in separate
# folders are merged.
def latest_by(rows, key):
    latest = {}
    for r in rows:
        latest[(r[key], r["seed"])] = r
    groups = {}
    for (value, _), r in latest.items():
        groups.setdefault(value, []).append(r)
    return groups


def load_probabilities(path):
    with np.load(path) as saved:
        return saved["labels"].astype(int), saved["probabilities"]


# Paired DeLong test of `code` against `reference`, seed by seed. Seeds missing
# on either side are skipped.
def delong_rows(code, group, reference, reference_group):
    by_seed = {r["seed"]: r for r in reference_group}
    table = []
    for r in sorted(group, key=lambda r: r["seed"]):
        other = by_seed.get(r["seed"])
        if other is None or not r["test_probs"].exists() or not other["test_probs"].exists():
            continue
        labels, a = load_probabilities(r["test_probs"])
        other_labels, b = load_probabilities(other["test_probs"])
        row = {"Kịch bản": code, "So với": reference, "Seed": r["seed"]}
        if np.array_equal(labels, other_labels):
            auc_a, auc_b, z, p = delong(labels, a, b)
            row.update({"AUC kịch bản": round(float(auc_a), 4), "AUC so với": round(float(auc_b), 4),
                        "Δ AUC": round(float(auc_a - auc_b), 4), "z": round(float(z), 2),
                        "p": float(f"{p:.3g}"), "Ghi chú": ""})
        else:
            row["Ghi chú"] = "Nhãn test khác nhau (tập hoặc thứ tự test khác), không so được"
        table.append(row)
    return table


# Bang_chinh: one row per scenario of scripts/scenarios.py, in its order.
def main_table(by_code, comparisons):
    table = []
    for code, spec in SCENARIOS.items():
        group = by_code.get(code, [])
        row = {"Mã": code, "Nhóm": spec["group"], "Mô tả": spec["description"], "Số seed": len(group)}
        if not group:
            row["Trạng thái"] = "Chưa chạy"
            table.append(row)
            continue
        summary = summarize(group, {})
        row["Trạng thái"] = "Đủ 5 seed" if len(group) >= 5 else "Chưa đủ 5 seed"
        row["Seeds"] = summary["Seeds"]
        row["User rule"] = summary["User rule"]
        for _, header in METRICS:
            row[f"{header} test"] = summary[f"{header} test"]
        for header in ("AUROC (AUC)", "AUPRC"):
            row[f"{header} val"] = summary[f"{header} val"]
        auc = mean_std(r["test_auroc"] for r in group)[0]
        for reference in REFERENCES:
            if code == reference or reference not in by_code:
                continue
            reference_auc = mean_std(r["test_auroc"] for r in by_code[reference])[0]
            if auc is not None and reference_auc is not None:
                row[f"Δ AUC so với {reference}"] = round(auc - reference_auc, 4)
            p_values = [c["p"] for c in comparisons
                        if c["Kịch bản"] == code and c["So với"] == reference and "p" in c]
            if p_values:
                row[f"DeLong vs {reference}: seed p<0,05"] = f"{sum(p < 0.05 for p in p_values)}/{len(p_values)}"
                row[f"DeLong vs {reference}: p lớn nhất"] = max(p_values)
        row["Lần chạy"] = ", ".join(sorted({r["folder"] for r in group}))
        table.append(row)
    # The same columns in the same order for every row, scenarios not run included.
    columns = ["Mã", "Nhóm", "Mô tả", "Số seed", "Trạng thái", "Seeds", "User rule"]
    columns += [f"{header} test" for _, header in METRICS] + ["AUROC (AUC) val", "AUPRC val"]
    for reference in REFERENCES:
        columns += [f"Δ AUC so với {reference}", f"DeLong vs {reference}: seed p<0,05",
                    f"DeLong vs {reference}: p lớn nhất"]
    columns.append("Lần chạy")
    return [{column: row.get(column) for column in columns} for row in table]


# Kich_ban: the scenario catalog, and how many seeds each one has so far.
def catalog_table(by_code):
    return [{"Mã": code, "Nhóm": spec["group"], "Mô tả": spec["description"], "Câu hỏi": spec["question"],
             "Script": spec["script"], "Tham số": " ".join(spec["args"]),
             "File graph": hypergraph_file(spec) or "", "User rule cần": spec["user_rule"] or "",
             "Số seed đã chạy": len(by_code.get(code, []))}
            for code, spec in SCENARIOS.items()]


GUIDE = (
    ("Bang_chinh", "Mỗi dòng là một kịch bản của scripts/scenarios.py (mã kịch bản = --tag của lần chạy). "
                   "Với mỗi seed lấy lần chạy mới nhất, nên các seed chạy ở nhiều thư mục được gộp lại."),
    ("Kich_ban", "Danh mục kịch bản: câu hỏi cần trả lời, script, tham số dòng lệnh, file graph, User rule cần. "
                 "Siêu tham số không ghi ở đây lấy mặc định của DEFAULT_SETTINGS trong src/8_train.py; "
                 "giá trị thật của từng lần chạy nằm ở Lan_chay."),
    ("Lan_chay", "Mỗi dòng là một thư mục result/<ngày_giờ> (một lần gọi run_all.sh): git commit, máy, GPU, "
                 "mọi siêu tham số, mean ± std trên các seed của nó."),
    ("Theo_seed", "Từng seed, số chưa làm tròn thành mean ± std."),
    ("DeLong", "Kiểm định DeLong ghép cặp trên test, từng seed, mỗi kịch bản so với M0 và với M-any. "
               "DeLong coi dự đoán là cố định nên không tính dao động giữa các seed: đọc kèm mean ± std."),
    ("Theo_cau_hinh", "Mỗi dòng là một cấu hình (tên run bỏ phần seed), gồm cả các lần chạy trước khi có mã kịch bản."),
    ("User rule", "any: User hyperedge nối mọi lượt đăng ký của học viên, kể cả khoá bắt đầu sau (dùng thông tin "
                  "tương lai, như các baseline đã công bố). temporal: chỉ các khoá đã bắt đầu, kèm --causal. "
                  "Đọc từ file graph lúc train; trống = lần chạy cũ không ghi."),
    ("Δ AUC so với M0 / M-any", "AUC test trung bình của kịch bản trừ AUC test trung bình của M0 (hoặc M-any)."),
    ("AUROC (AUC)", "Diện tích dưới đường ROC. AUC và AUROC là cùng một chỉ số. Không phụ thuộc ngưỡng."),
    ("AUPRC", "Diện tích dưới đường Precision–Recall của lớp bỏ học (average precision)."),
    ("ACC", "Tỉ lệ dự đoán đúng tại ngưỡng t*. Lớp bỏ học chiếm 75,8%, nên đoán 'ai cũng bỏ' đã đạt ACC 0,758."),
    ("Precision, Recall, F1", "Của lớp bỏ học (nhãn 1), tại ngưỡng t*."),
    ("Specificity", "Recall của lớp không bỏ học: tỉ lệ học viên không bỏ được nhận ra đúng."),
    ("Macro-F1", "Trung bình F1 của hai lớp."),
    ("Ngưỡng t*", "Chọn trên toàn bộ validation để F1 lớp bỏ học cao nhất, rồi dùng nguyên cho test."),
    ("kept_*", "HSL: tỉ lệ membership được giữ lại theo loại hyperedge, ở best epoch."),
    ("w_*", "Trọng số family học được (--family-weights), ở best epoch. Chỉ tỉ lệ giữa các w có nghĩa."),
    ("alpha_*", "Trung bình trọng số từng hyperedge α_e theo loại (--edge-weights). 1 = chưa khác mặc định."),
    ("Chọn mô hình", "So sánh và chọn bằng cột val; cột test chỉ để báo cáo."),
)


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, table in sheets:
        sheet = workbook.create_sheet(name)
        columns = table_columns(table)
        sheet.append(columns)
        for row in table:
            sheet.append([row.get(column) for column in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        sheet.freeze_panes = "B2" if name != "Giai_thich" else None
        for index, column in enumerate(columns, 1):
            width = max([len(str(column))] + [len(str(row.get(column) or "")) for row in table[:200]])
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = min(max(10, width + 2), 60)
    workbook.save(path)


# Write to a temporary file, then rename: a reader never sees half a workbook.
def save_workbook(path, sheets):
    temporary = path.with_name(f".{path.stem}.{os.getpid()}.xlsx")
    write_xlsx(temporary, sheets)
    try:
        os.replace(temporary, path)
    except PermissionError:
        temporary.unlink(missing_ok=True)
        print(f"{path} is open in another program (Excel?); close it and run "
              "python scripts/summarize_results.py again")


def write_markdown(path, table):
    columns = ["Mã", "Mô tả", "Số seed", "User rule", "AUROC (AUC) test", "AUPRC test", "F1 test",
               "Macro-F1 test", "Δ AUC so với M0", "DeLong vs M0: seed p<0,05"]
    lines = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    for row in table:
        lines.append("| " + " | ".join(str(row.get(column, "")) for column in columns) + " |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(result_dir):
    # A Windows console (cp1252) cannot print Vietnamese; replace instead of failing.
    sys.stdout.reconfigure(errors="replace")
    result_dir = Path(result_dir)
    folders = sorted((d for d in result_dir.iterdir() if (d / "reports").is_dir()), key=folder_time)
    rows = []
    for folder in folders:
        rows.extend(seed_rows(folder))
    if not rows:
        raise SystemExit(f"No run with reports under {result_dir}")

    by_folder = {}
    for r in rows:
        by_folder.setdefault(r["folder"], []).append(r)
    runs = []
    for number, (folder, group) in enumerate(by_folder.items(), 1):
        runs.append(summarize(group, {"STT": number, "Lần chạy": folder,
                                      "Ngày": group[0]["started"].strftime("%d/%m/%Y %H:%M"),
                                      "Cấu hình": group[-1]["scenario"], "Ghi chú": group[-1]["note"],
                                      "Git": group[-1]["git"], "Máy": group[-1]["host"],
                                      "GPU": group[-1]["gpu"], "Torch": group[-1]["torch"]}))

    configurations = [summarize(group, {"Cấu hình": name,
                                        "Lần chạy": ", ".join(sorted({r["folder"] for r in group}))})
                      for name, group in latest_by(rows, "scenario").items()]
    configurations.sort(key=lambda row: -row["_sort"])

    by_tag = latest_by(rows, "tag")
    by_code = {code: by_tag[code] for code in SCENARIOS if code in by_tag}
    comparisons = []
    for code, group in by_code.items():
        for reference in REFERENCES:
            if code != reference and reference in by_code:
                comparisons.extend(delong_rows(code, group, reference, by_code[reference]))
    main_rows = main_table(by_code, comparisons)

    seeds = seed_table(rows)
    write_csv(result_dir / "tong_hop_bang_chinh.csv", main_rows)
    write_csv(result_dir / "tong_hop_lan_chay.csv", runs)
    write_csv(result_dir / "tong_hop_cau_hinh.csv", configurations)
    write_csv(result_dir / "tong_hop_theo_seed.csv", seeds)
    write_markdown(result_dir / "tong_hop_ket_qua.md", main_rows)
    try:
        guide = [{"Mục": name, "Giải thích": meaning} for name, meaning in GUIDE]
        save_workbook(result_dir / WORKBOOK,
                      [("Bang_chinh", main_rows), ("Kich_ban", catalog_table(by_code)), ("Lan_chay", runs),
                       ("Theo_seed", seeds), ("DeLong", comparisons), ("Theo_cau_hinh", configurations),
                       ("Giai_thich", guide)])
    except ImportError:
        print("openpyxl not installed; wrote CSV and Markdown only")
    print(f"{len(runs)} run folders, {len(by_code)}/{len(SCENARIOS)} scenarios, {len(configurations)} "
          f"configurations, {len(seeds)} seeds -> {result_dir / WORKBOOK}")
    for row in main_rows:
        print(f"  {row['Mã']:<7} n={row['Số seed']}  test AUROC {row['AUROC (AUC) test'] or '-':<17}  "
              f"dAUC vs M0 {row['Δ AUC so với M0'] or '-'}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Rebuild the experiment workbook from every run folder in result/.")
    parser.add_argument("--result-dir", type=Path, default=ROOT / "result")
    main(parser.parse_args().result_dir)
