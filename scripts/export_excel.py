# Add one scripts/run_all.sh run to the experiment log workbook.
#
#   python scripts/export_excel.py result/<dd-mm-yyyy_HH-MM> [--note "..."]
#
# Reads <run>/reports/*_train.json and *_test.json (copied there by
# collect_results.py) and <run>/run_info.txt, then writes into
# docs/ket_qua_thi_nghiem.xlsx (created if missing):
#   Ket_qua        one row per run: STT, date, configuration, test metrics as
#                  "mean ± std" (Excel formulas over Du_lieu_seed)
#   Sieu_tham_so   one row per run, same STT: every setting of the run
#   Du_lieu_seed   one row per seed: the numbers the formulas above read
# Running it again on the same run folder replaces that run's rows.

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
DEFAULT_XLSX = ROOT / "docs" / "ket_qua_thi_nghiem.xlsx"
METRICS = ("auc", "auprc", "f1", "precision", "recall")
LAST_ROW = 5000  # formulas in Ket_qua read Du_lieu_seed rows 2..LAST_ROW

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLUE = "0000FF"
GREY = "808080"

RESULT_HEADERS = ("STT", "Ngày chạy", "Cấu hình", "AUC", "AUPRC", "F1", "Precision", "Recall",
                  "Ghi chú", "Mã lần chạy")
RESULT_WIDTHS = {1: 6, 2: 14, 3: 34, 4: 17, 5: 17, 6: 17, 7: 17, 8: 17, 9: 60, 10: 18}
FIRST_METRIC_COLUMN = 4
RESULT_ID_COLUMN = len(RESULT_HEADERS)

# (header, key in the train settings or in extra_settings)
PARAMETERS = (
    ("Seeds", "seeds"),
    ("Tập đặc trưng", "feature_set"),
    ("Dùng HSL", "hsl"),
    ("Me: lấy mẫu siêu cạnh", "edge_sampling"),
    ("Mv: lấy mẫu liên thuộc", "node_sampling"),
    ("ΔH: số node thêm mỗi cạnh", "add_per_edge"),
    ("Chiều ẩn", "hidden_dim"),
    ("Dropout", "dropout"),
    ("Learning rate", "learning_rate"),
    ("Weight decay", "weight_decay"),
    ("Epoch tối đa", "epochs"),
    ("Đánh giá mỗi (epoch)", "eval_every"),
    ("Patience", "patience"),
    ("λ tương phản", "lambda_cl"),
    ("τ tương phản", "temperature"),
    ("Số node neo", "contrastive_anchors"),
    ("Số mẫu âm K", "contrastive_neighbors"),
    ("τ Gumbel", "gumbel_temperature"),
    ("Chiều MLP chấm điểm", "scorer_dim"),
    ("k (behavioral)", "k"),
    ("k_max (ứng viên ΔH)", "k_max"),
    ("Cửa sổ quan sát (ngày)", "observation_days"),
    ("Tỉ lệ train/val", "train_ratio"),
    ("Seed chia dữ liệu", "split_seed"),
    ("Giới hạn target val", "validation_limit"),
    ("Batch đánh giá", "eval_batch_size"),
)
PARAMETER_HEADERS = ("STT", "Ngày chạy", "Cấu hình") + tuple(h for h, _ in PARAMETERS) + ("Mã lần chạy",)
PARAMETER_ID_COLUMN = len(PARAMETER_HEADERS)

# (header, key in the seed row, number format)
SEED_COLUMNS = (
    ("Mã lần chạy", "run_id", None),
    ("Seed", "seed", "0"),
    ("Tên run", "run_name", None),
    ("Best epoch", "best_epoch", "0"),
    ("Số epoch đã chạy", "epochs_run", "0"),
    *[(f"{name.upper()} val", f"val_{name}", "0.0000") for name in METRICS],
    *[(f"{name.upper()} test", f"test_{name}", "0.0000") for name in METRICS],
    ("Giữ course", "kept_course", "0.000"),
    ("Giữ object", "kept_object", "0.000"),
    ("Giữ behavioral", "kept_behavioral", "0.000"),
    ("ΔH được giữ", "added", "#,##0"),
    ("Thời gian (giây)", "duration_sec", "#,##0"),
)
SEED_LETTER = {key: get_column_letter(index) for index, (_, key, _) in enumerate(SEED_COLUMNS, 1)}


def style_header(sheet, headers, widths=None):
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        width = (widths or {}).get(column, max(10, min(len(str(header)) + 2, 20)))
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 36
    sheet.freeze_panes = "D2" if headers[:3] == ("STT", "Ngày chạy", "Cấu hình") else "B2"


def put(sheet, row, column, value, number_format=None, *, color=None, align=None):
    cell = sheet.cell(row=row, column=column, value=value)
    cell.font = Font(name=FONT, color=color)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if number_format:
        cell.number_format = number_format
    return cell


def new_workbook():
    workbook = Workbook()
    results = workbook.active
    results.title = "Ket_qua"
    style_header(results, RESULT_HEADERS, RESULT_WIDTHS)
    style_header(workbook.create_sheet("Sieu_tham_so"), PARAMETER_HEADERS, {1: 6, 2: 14, 3: 34, 4: 20})
    style_header(workbook.create_sheet("Du_lieu_seed"), [h for h, _, _ in SEED_COLUMNS], {1: 18, 3: 22})
    add_guide_sheet(workbook.create_sheet("Huong_dan"))
    workbook.calculation.fullCalcOnLoad = True
    return workbook


def add_guide_sheet(sheet):
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 100
    lines = (
        ("Thêm một lần chạy", "python scripts/export_excel.py result/<dd-mm-yyyy_HH-MM> --note \"...\""),
        ("", "Chạy lại trên cùng thư mục sẽ thay các dòng cũ của lần chạy đó (giữ nguyên STT)."),
        ("Ket_qua", "Chỉ số trên tập test, dạng trung bình ± độ lệch chuẩn mẫu (n − 1) qua các seed. "
                    "Công thức đọc từ Du_lieu_seed."),
        ("Sieu_tham_so", "Cùng STT với Ket_qua. Toàn bộ settings của 8_train.py, cộng τ Gumbel, chiều MLP "
                         "chấm điểm (6_hsl.py), k, k_max (hypergraph.npz), cấu hình chia dữ liệu (0_config.py)."),
        ("Du_lieu_seed", "Số liệu từng seed. 'Giữ ...' và 'ΔH được giữ' lấy ở best epoch."),
        ("Chữ xanh dương", "Số nhập tay, chép từ tài liệu, không phải công thức."),
        ("Cột Ghi chú", "Em tự sửa được; script chỉ ghi đè khi truyền --note."),
    )
    for row, (key, text) in enumerate(lines, 1):
        put(sheet, row, 1, key).font = Font(name=FONT, bold=True)
        put(sheet, row, 2, text)


# ---------------------------------------------------------------------------
# Reading one run folder
# ---------------------------------------------------------------------------

def read_run_info(run_dir):
    info = {}
    path = run_dir / "run_info.txt"
    if path.exists():
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            key, _, value = line.partition(":")
            info[key.strip()] = value.strip()
    return info


def read_durations(run_dir):
    durations = {}
    path = run_dir / "manifest.tsv"
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines()[1:]:
            seed, _, duration, *_ = line.split("\t")
            if duration:
                durations[int(seed)] = float(duration)
    return durations


def read_seed_rows(run_dir, run_id):
    durations = read_durations(run_dir)
    rows = []
    settings = None
    for train_path in sorted((run_dir / "reports").glob("*_train.json")):
        test_path = train_path.with_name(train_path.name.replace("_train.json", "_test.json"))
        if not test_path.exists():
            print(f"  skip {train_path.name}: no test report")
            continue
        train = json.loads(train_path.read_text(encoding="utf-8"))
        test = json.loads(test_path.read_text(encoding="utf-8"))["test"]
        settings = train["settings"]
        best = next((r for r in train["history"] if r["epoch"] == train["best_epoch"]), {})
        row = {
            "run_id": run_id,
            "seed": train["seed"],
            "run_name": train_path.name.removesuffix("_train.json"),
            "best_epoch": train["best_epoch"],
            "epochs_run": len(train["history"]),
            "duration_sec": durations.get(train["seed"]),
        }
        for name in METRICS:
            row[f"val_{name}"] = train["best_validation"][name]
            row[f"test_{name}"] = test[name]
        for key in ("kept_course", "kept_object", "kept_behavioral", "added"):
            row[key] = best.get(key)
        rows.append(row)
    return sorted(rows, key=lambda row: row["seed"]), settings


# Settings that are not in the train report: code defaults and graph parameters.
def extra_settings(processed_dir):
    from importlib import import_module
    import inspect
    config = import_module("0_config")
    extra = {
        "observation_days": config.OBSERVATION_DAYS,
        "train_ratio": config.TRAIN_RATIO,
        "split_seed": config.SPLIT_SEED,
    }
    try:
        defaults = inspect.signature(import_module("6_hsl").StructureLearner.__init__).parameters
        extra["gumbel_temperature"] = defaults["temperature"].default
        extra["scorer_dim"] = defaults["scorer_dim"].default
    except ImportError:
        pass
    bundle_path = Path(processed_dir) / "hypergraph.npz"
    if bundle_path.exists():
        import numpy as np
        with np.load(bundle_path) as bundle:
            extra["k"] = int(bundle["k"])
            extra["k_max"] = int(bundle["train_neighbors"].shape[1])
    return extra


# Readable name of a configuration, e.g. "HSL, bỏ ΔH (feature: full)".
def describe(settings):
    if not settings["hsl"]:
        name = "HGNN (không HSL)"
    else:
        removed = []
        if not settings["edge_sampling"]:
            removed.append("Me")
        if not settings["node_sampling"]:
            removed.append("Mv")
        if settings["add_per_edge"] == 0:
            removed.append("ΔH")
        if settings["lambda_cl"] == 0:
            removed.append("tương phản")
        name = "HSL đầy đủ" if not removed else "HSL, bỏ " + ", ".join(removed)
    return f"{name} (feature: {settings['feature_set']})"


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------

# "mean ± std" of one Du_lieu_seed column over the seeds of this run.
# FIXED follows the decimal separator of the Excel locale.
def metric_formula(row, key):
    ids = f"Du_lieu_seed!$A$2:$A${LAST_ROW}"
    values = f"Du_lieu_seed!${SEED_LETTER[key]}$2:${SEED_LETTER[key]}${LAST_ROW}"
    run = f"${get_column_letter(RESULT_ID_COLUMN)}{row}"
    count = f"COUNTIFS({ids},{run})"
    mean = f"AVERAGEIFS({values},{ids},{run})"
    std = f"SQRT(SUMPRODUCT(({ids}={run})*({values}-{mean})^2)/({count}-1))"
    return (f'=IFERROR(IF({count}=0,"",IF({count}=1,FIXED({mean},3),'
            f'FIXED({mean},3)&" ± "&FIXED({std},3))),"")')


def find_row(sheet, id_column, run_id):
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=id_column).value == run_id:
            return row
    return None


def next_number(sheet):
    numbers = [sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)]
    return max([n for n in numbers if isinstance(n, int)], default=0) + 1


def write_result_row(sheet, row, number, date, config_name, run_id, note):
    put(sheet, row, 1, number, "0", align="center")
    put(sheet, row, 2, date, align="center")
    put(sheet, row, 3, config_name)
    for offset, name in enumerate(METRICS):
        put(sheet, row, FIRST_METRIC_COLUMN + offset, metric_formula(row, f"test_{name}"), align="center")
    put(sheet, row, RESULT_ID_COLUMN - 1, note)
    put(sheet, row, RESULT_ID_COLUMN, run_id, color=GREY)


def write_parameter_row(sheet, row, number, date, config_name, run_id, values):
    put(sheet, row, 1, number, "0", align="center")
    put(sheet, row, 2, date, align="center")
    put(sheet, row, 3, config_name)
    for offset, (_, key) in enumerate(PARAMETERS):
        value = values.get(key)
        if isinstance(value, bool):
            value = "Có" if value else "Không"
        put(sheet, row, 4 + offset, value, align="center")
    put(sheet, row, PARAMETER_ID_COLUMN, run_id, color=GREY)


def export(run_dir, xlsx_path, *, run_id=None, note=None, processed_dir=None):
    run_dir = Path(run_dir)
    run_id = run_id or run_dir.name
    rows, settings = read_seed_rows(run_dir, run_id)
    if not rows:
        sys.exit(f"No seed with both train and test reports in {run_dir / 'reports'}")

    xlsx_path = Path(xlsx_path)
    workbook = load_workbook(xlsx_path) if xlsx_path.exists() else new_workbook()
    results = workbook["Ket_qua"]
    parameters = workbook["Sieu_tham_so"]
    seeds = workbook["Du_lieu_seed"]

    for row in range(seeds.max_row, 1, -1):
        if seeds.cell(row=row, column=1).value == run_id:
            seeds.delete_rows(row)
    for seed_row in rows:
        target = seeds.max_row + 1
        for column, (_, key, number_format) in enumerate(SEED_COLUMNS, 1):
            put(seeds, target, column, seed_row.get(key), number_format)

    info = read_run_info(run_dir)
    date = info.get("started", "").split(" ")[0] or None
    config_name = describe(settings)
    result_row = find_row(results, RESULT_ID_COLUMN, run_id)
    if result_row is None:
        result_row, number = results.max_row + 1, next_number(results)
        old_note = None
    else:
        number = results.cell(row=result_row, column=1).value
        old_note = results.cell(row=result_row, column=RESULT_ID_COLUMN - 1).value
    write_result_row(results, result_row, number, date, config_name, run_id,
                     note or old_note or info.get("command"))

    values = {**settings, **extra_settings(processed_dir or ROOT / "data" / "processed" / "simple")}
    values["seeds"] = " ".join(str(seed_row["seed"]) for seed_row in rows)
    parameter_row = find_row(parameters, PARAMETER_ID_COLUMN, run_id) or parameters.max_row + 1
    write_parameter_row(parameters, parameter_row, number, date, config_name, run_id, values)

    workbook.calculation.fullCalcOnLoad = True
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)
    print(f"Saved {xlsx_path}: STT {number}, {rows[0]['run_name']}, seeds {values['seeds']}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Add one run folder to the experiment log workbook.")
    parser.add_argument("run_dir", type=Path)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--run-id", help="default: the run folder name (use lan_1 to fill run 1)")
    parser.add_argument("--note", help="text for the Ghi chú column")
    parser.add_argument("--processed-dir", type=Path, help="where hypergraph.npz is (for k, k_max)")
    arguments = parser.parse_args()
    export(arguments.run_dir, arguments.xlsx, run_id=arguments.run_id,
           note=arguments.note, processed_dir=arguments.processed_dir)
